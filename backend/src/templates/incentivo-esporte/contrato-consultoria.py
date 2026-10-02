# Gera contrato-consultoria.docx (template docxtemplater) — rodar de novo
# depois de mudar o texto:
#   python3 backend/src/templates/incentivo-esporte/contrato-consultoria.py
# Texto espelha o modelo "Contrato de Consultoria — Lei de Incentivo ao
# Esporte de Curitiba" (Claude Docs). Blocos condicionais: {#pf}/{#pj}/
# {#menor}/{#pagamento_unico}/{#pagamento_proporcional} sozinhos no
# parágrafo (paragraphLoop remove o parágrafo da tag). Cada placeholder vai
# num run inteiro pra não ser quebrado.
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.5)
estilo = doc.styles['Normal']
estilo.font.name = 'Arial'
estilo.font.size = Pt(11)
estilo.paragraph_format.space_after = Pt(6)


def par(texto, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, tamanho=None):
    """**negrito** vira run em negrito; o resto, runs normais."""
    p = doc.add_paragraph()
    p.alignment = alinhamento
    for i, pedaco in enumerate(re.split(r'\*\*', texto)):
        if not pedaco:
            continue
        run = p.add_run(pedaco)
        run.bold = i % 2 == 1
        if tamanho:
            run.font.size = Pt(tamanho)
    return p


def tag(t):
    doc.add_paragraph(t)


def titulo(t):
    par(f'**{t}**', WD_ALIGN_PARAGRAPH.LEFT)


par('**CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE CONSULTORIA — PROGRAMA MUNICIPAL DE INCENTIVO AO ESPORTE DE CURITIBA**', WD_ALIGN_PARAGRAPH.CENTER, 13)

titulo('DAS PARTES')
par('**CONTRATADO:** {contratado_nome}, {contratado_nacionalidade}, {contratado_estado_civil}, {contratado_profissao}, portador do RG nº {contratado_rg} e do CPF nº {contratado_cpf}, residente em {contratado_endereco}, e-mail {contratado_email}, telefone {contratado_telefone}.')
tag('{#pf}')
par('**CONTRATANTE:** {contratante_nome}, portador(a) do RG nº {contratante_rg} e do CPF nº {contratante_cpf}, nascido(a) em {contratante_nascimento}, residente em {contratante_endereco}.')
tag('{/pf}')
tag('{#menor}')
par('Por ser menor de 18 anos, o CONTRATANTE é representado(a) por seu(sua) responsável legal {responsavel_nome}, RG nº {responsavel_rg}, CPF nº {responsavel_cpf}, que assina este contrato e responde por todas as obrigações aqui assumidas.')
tag('{/menor}')
tag('{#pj}')
par('**CONTRATANTE:** {pj_razao_social}, entidade sem fins lucrativos inscrita no CNPJ sob o nº {pj_cnpj}, com sede em {pj_sede}, neste ato representada por seu(sua) {pj_cargo_representante} {pj_representante_nome}, RG nº {pj_representante_rg}, CPF nº {pj_representante_cpf}, conforme {pj_documento_representacao}.')
tag('{/pj}')
par('As partes acima identificadas têm entre si justo e contratado o que segue.')

titulo('CLÁUSULA 1ª — DO OBJETO')
par('1.1. O CONTRATADO prestará consultoria ao CONTRATANTE para o projeto {projeto_nome} no Programa Municipal de Incentivo ao Esporte de Curitiba (Decreto Municipal nº 1985/2025 e Resolução CIE nº 004/2026), com protocolo em outubro de 2026 e execução de 01/03 a 31/12/2027.')
par('1.2. A consultoria compreende três etapas:')
par('a) **Elaboração e protocolo:** análise de elegibilidade e classificação provável, lista e conferência dos documentos exigidos, preenchimento dos anexos padrão, redação do projeto e do plano de aplicação, e protocolo no Sistema Incentivo online até o fim do prazo oficial.')
par('b) **Acompanhamento:** retificação ou complementação de documentos no prazo do edital, eventuais recursos, aceite do Termo de Compromisso e orientação durante a execução sobre despesas permitidas, contrapartida social e divulgação obrigatória.')
par('c) **Prestação de contas:** organização dos comprovantes fiscais e protocolo das prestações de contas parcial (se exigida) e final.')
par('1.3. Não fazem parte do objeto: custas de certidões, reconhecimento de firma, taxas e emissão de documentos; representação judicial; e assessoria contábil ou jurídica.')

titulo('CLÁUSULA 2ª — DAS OBRIGAÇÕES DO CONTRATADO')
par('2.1. Executar as etapas da Cláusula 1ª com diligência, seguindo o Decreto, a Resolução e o Manual do Beneficiário vigentes.')
par('2.2. Informar ao CONTRATANTE, por escrito (e-mail ou WhatsApp), a lista de documentos necessários e os prazos de cada etapa, com antecedência razoável.')
par('2.3. Submeter ao CONTRATANTE, antes do protocolo, a versão final do projeto e do plano de aplicação, e só protocolar após a aprovação dele.')
par('2.4. Orientar o CONTRATANTE a usar o recurso somente nas despesas autorizadas (Anexos VI, VII ou VIII da Resolução) e alertá-lo sobre as vedações do art. 16 do Decreto.')
par('2.5. Manter o CONTRATANTE informado sobre publicações, resultados e pendências do projeto.')

titulo('CLÁUSULA 3ª — DAS OBRIGAÇÕES DO CONTRATANTE')
par('3.1. Fornecer, nos prazos combinados, documentos verdadeiros, legíveis e válidos, e responder pela veracidade de todas as informações prestadas (art. 19, §1º da Resolução).')
par('3.2. Colher as assinaturas exigidas, inclusive de terceiros (atletas, responsáveis legais, proprietário do imóvel, responsável técnico), com firma reconhecida ou assinatura digital quando o edital exigir.')
par('3.3. Após a aprovação: executar o projeto, cumprir a contrapartida social (no mínimo 4 campanhas) e a divulgação obrigatória do Município e do Programa (art. 37 da Resolução).')
par('3.4. Usar o recurso apenas nas despesas autorizadas, guardar todas as notas fiscais e entregá-las ao CONTRATADO em até {prazo_notas_dias} dias da despesa, para a prestação de contas.')
par('3.5. Comunicar ao CONTRATADO, em até {prazo_notificacao_dias} dias úteis, qualquer notificação recebida da Secretaria Municipal do Esporte, Lazer e Juventude.')
tag('{#pj}')
par('3.6. Não contratar, com recurso do incentivo, fornecedor ou prestador que seja membro da entidade ou parente até 2º grau de membro (art. 16, §1º do Decreto).')
tag('{/pj}')

titulo('CLÁUSULA 4ª — DO ACESSO AO SISTEMA INCENTIVO ONLINE')
# PF: login do e-Cidadão pelo CPF do proponente. PJ: login criado para o
# CNPJ da entidade, vinculado ao e-mail institucional dela.
tag('{#pf}')
par('4.1. O protocolo é feito com login e senha individuais do e-Cidadão, pelo CPF do próprio proponente (arts. 13 e 19 da Resolução).')
par('4.2. O CONTRATANTE (ou o representante legal que assina este contrato) **autoriza expressamente** o CONTRATADO a acessar o Sistema Incentivo online com essas credenciais, exclusivamente para executar o objeto deste contrato.')
par('4.3. O CONTRATANTE declara estar ciente de que todo protocolo feito com suas credenciais é feito em seu nome e sob sua responsabilidade perante a Prefeitura, e que conferiu e aprovou o conteúdo antes do envio (cláusula 2.3).')
tag('{/pf}')
tag('{#pj}')
par('4.1. O acesso ao Sistema Incentivo online é feito por login cadastrado para o CNPJ do CONTRATANTE, vinculado ao e-mail institucional da entidade ({pj_email_institucional}), que o CONTRATANTE fornecerá ao CONTRATADO e manterá ativo durante a vigência deste contrato.')
par('4.2. O CONTRATANTE, por seu representante legal, **autoriza expressamente** o CONTRATADO a criar, se ainda não existir, e a acessar esse login, exclusivamente para executar o objeto deste contrato.')
par('4.3. O CONTRATANTE declara estar ciente de que todo protocolo feito com esse login é feito em nome da entidade e sob sua responsabilidade perante a Prefeitura, e que conferiu e aprovou o conteúdo antes do envio (cláusula 2.3).')
tag('{/pj}')
par('4.4. O CONTRATADO guardará as credenciais com sigilo, sem repassá-las a terceiros, e não as usará para nada além do objeto. Ao fim do contrato, o CONTRATANTE pode alterar sua senha, e o CONTRATADO apagará as credenciais que mantiver.')

titulo('CLÁUSULA 5ª — DA REMUNERAÇÃO')
par('5.1. Pelos serviços, o CONTRATANTE pagará ao CONTRATADO **{percentual} ({percentual_extenso}) do valor total aprovado** para o projeto no Programa.')
par('5.2. **Sem aprovação, nada é devido.** Se o projeto for indeferido, inabilitado ou não contemplado por falta de orçamento, o CONTRATANTE não paga nenhum valor.')
par('5.3. **Origem do pagamento:** a remuneração será paga com recursos próprios do CONTRATANTE, **nunca com o recurso do incentivo**. Consultoria não é despesa autorizada pelo Programa (art. 15 do Decreto; Anexos VI a VIII da Resolução), e o uso da verba para esse fim sujeita o beneficiário a devolução e penalidades.')
tag('{#pagamento_unico}')
par('5.4. **Forma e prazo:** em parcela única, até {prazo_pagamento_dias} dias após o recebimento da primeira parcela do incentivo, via {forma_pagamento}.')
tag('{/pagamento_unico}')
tag('{#pagamento_proporcional}')
par('5.4. **Forma e prazo:** proporcionalmente, {percentual} de cada parcela do incentivo, até {prazo_pagamento_dias} dias após cada recebimento, via {forma_pagamento}.')
tag('{/pagamento_proporcional}')
par('5.5. Se o valor aprovado for reduzido, suspenso ou tiver de ser devolvido por fato imputável ao CONTRATANTE, a remuneração continua calculada sobre o valor aprovado. Se a redução decorrer de erro do CONTRATADO, a remuneração será recalculada sobre o valor efetivamente recebido.')
par('5.6. Atraso no pagamento gera multa de {multa_percentual} e juros de 1% ao mês, mais correção pelo {indice_correcao}.')
par('5.7. O CONTRATADO emitirá {comprovante_pagamento} de cada pagamento.')

titulo('CLÁUSULA 6ª — DA AUSÊNCIA DE GARANTIA E DOS LIMITES DE RESPONSABILIDADE')
par('6.1. A consultoria é obrigação de meio. O CONTRATADO **não garante** a aprovação do projeto nem o valor a ser concedido, que dependem da análise do Comitê de Avaliação, da classificação do proponente e da disponibilidade orçamentária do Programa (arts. 6º e 10 da Resolução).')
par('6.2. O CONTRATADO não responde por indeferimento, inabilitação, glosa ou devolução causados por: documento falso, vencido ou entregue fora do prazo pelo CONTRATANTE; despesa feita sem consulta ao CONTRATADO ou contra a orientação dele; descumprimento de contrapartida ou divulgação; ou mudança nas normas do Programa.')
par('6.3. O CONTRATADO responde pelos prejuízos causados por sua culpa, como perda de prazo de protocolo por ele assumido com toda a documentação entregue a tempo, limitada a indenização ao valor da remuneração prevista na Cláusula 5ª.')

titulo('CLÁUSULA 7ª — DA CONFIDENCIALIDADE E PROTEÇÃO DE DADOS')
par('7.1. O CONTRATADO tratará os dados pessoais recebidos (documentos, endereços, certidões, dados de menores e de terceiros) somente para executar este contrato, conforme a Lei nº 13.709/2018 (LGPD).')
par('7.2. O CONTRATANTE declara ter autorização dos terceiros cujos dados entrega ao CONTRATADO (atletas, responsáveis legais, colaboradores, proprietário do imóvel).')
par('7.3. Os dados ficarão guardados pelo prazo necessário à prestação de contas e a eventual fiscalização do Programa, e depois serão eliminados, salvo obrigação legal de guarda.')
par('7.4. As partes manterão sigilo sobre as informações trocadas, inclusive após o fim do contrato.')

titulo('CLÁUSULA 8ª — DA VIGÊNCIA, RESCISÃO E FORO')
par('8.1. Este contrato vigora da assinatura até a aprovação da prestação de contas final do projeto, ou até a publicação do resultado, se o projeto não for aprovado.')
par('8.2. Qualquer parte pode rescindir o contrato por aviso escrito com {aviso_rescisao_dias} dias de antecedência.')
par('8.3. Se o CONTRATANTE rescindir **após a aprovação do projeto**, sem culpa do CONTRATADO, a remuneração da Cláusula 5ª continua integralmente devida. Se rescindir antes da aprovação e o projeto protocolado pelo CONTRATADO vier a ser aprovado, idem.')
par('8.4. Se o CONTRATADO rescindir sem justa causa após a aprovação, a remuneração será devida apenas pelas etapas cumpridas, com estes pesos: {peso_elaboracao} elaboração e protocolo, {peso_acompanhamento} acompanhamento, {peso_prestacao} prestação de contas.')
par('8.5. Fica eleito o foro da Comarca de Curitiba/PR para dirimir questões deste contrato.')

par('E, por estarem de acordo, as partes assinam este contrato em 2 (duas) vias de igual teor, ou eletronicamente (gov.br ou certificado digital), na presença de 2 (duas) testemunhas.')
par('Curitiba, {data_extenso}.', WD_ALIGN_PARAGRAPH.RIGHT)

LINHA = '_' * 45
doc.add_paragraph()
par(LINHA, WD_ALIGN_PARAGRAPH.CENTER)
par('**CONTRATADO:** {contratado_nome} — CPF {contratado_cpf}', WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
par(LINHA, WD_ALIGN_PARAGRAPH.CENTER)
tag('{#pf}')
par('**CONTRATANTE:** {contratante_nome} — CPF {contratante_cpf}', WD_ALIGN_PARAGRAPH.CENTER)
tag('{/pf}')
tag('{#pj}')
par('**CONTRATANTE:** {pj_razao_social} — CNPJ {pj_cnpj}', WD_ALIGN_PARAGRAPH.CENTER)
par('p.p. {pj_representante_nome} — CPF {pj_representante_cpf}', WD_ALIGN_PARAGRAPH.CENTER)
tag('{/pj}')
tag('{#menor}')
doc.add_paragraph()
par(LINHA, WD_ALIGN_PARAGRAPH.CENTER)
par('**RESPONSÁVEL LEGAL:** {responsavel_nome} — CPF {responsavel_cpf}', WD_ALIGN_PARAGRAPH.CENTER)
tag('{/menor}')
doc.add_paragraph()
par('**Testemunhas:**', WD_ALIGN_PARAGRAPH.LEFT)
par('1. Nome: ______________________________ CPF: ________________ Assinatura: ________________', WD_ALIGN_PARAGRAPH.LEFT)
par('2. Nome: ______________________________ CPF: ________________ Assinatura: ________________', WD_ALIGN_PARAGRAPH.LEFT)

destino = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'contrato-consultoria.docx')
doc.save(destino)
print('ok', destino)
